from pathlib import Path
from openpyxl import Workbook
from openpyxl.styles import Font, Alignment
from db import get_flights

async def build_excel(start_date, end_date, out_path):
    flights=await get_flights(start_date,end_date)
    wb=Workbook()
    ws=wb.active
    ws.title="Flights"
    headers=[
        "ID","Дата","Рейс","Авиакомпания","Операция","Маршрут","Дней",
        "Гостиница экипажа","Гостиница пассажиров","Экипаж","Пассажиры",
        "Трансфер","Кем организован трансфер","Питание","Ресторан",
        "Доп. покупки","Что куплено","Кто обслуживал","Комментарий"
    ]
    ws.append(headers)
    for c in ws[1]:
        c.font=Font(bold=True); c.alignment=Alignment(horizontal="center")
    for f in flights:
        ws.append([
            f["id"],f["flight_date"],f["flight_no"],f["airline"],f["operation"],f["route"],f["stay_days"],
            f["crew_hotel"] or "",f["pax_hotel"] or "",f["crew_count"],f["pax_count"],
            "Да" if f["transfer"] else "Нет",f["transfer_provider"] or "",
            "Да" if f["catering"] else "Нет",f["restaurant"] or "",
            "Да" if f["extra_purchase"] else "Нет",f["extra_purchase_text"] or "",
            f["served_display"] or "",f["comment"] or ""
        ])
    ws.freeze_panes="A2"
    for col in ws.columns:
        ws.column_dimensions[col[0].column_letter].width=min(max(len(str(c.value or "")) for c in col)+2,40)
    sm=wb.create_sheet("Summary")
    sm["A1"],sm["B1"]="Период",f"{start_date} — {end_date}"
    sm["A2"],sm["B2"]="Всего рейсов",len(flights)
    sm["A3"],sm["B3"]="Прилёты",sum(1 for f in flights if f["operation"]=="Прилёт")
    sm["A4"],sm["B4"]="Вылеты",sum(1 for f in flights if f["operation"]=="Вылет")
    sm["A5"],sm["B5"]="Пассажиры",sum(int(f["pax_count"]) for f in flights)
    sm["A6"],sm["B6"]="Члены экипажа",sum(int(f["crew_count"]) for f in flights)
    Path(out_path).parent.mkdir(parents=True,exist_ok=True)
    wb.save(out_path)
    return out_path,len(flights)
